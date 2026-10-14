from pathlib import Path

import openpyxl

from vnhlfl_pipeline.build import LeagueBuilder, load
from vnhlfl_pipeline.mapping_sheet import main

from .workbooks import write_league


def test_sheet_lists_every_team_and_a_filled_in_gm_wins(tmp_path: Path) -> None:
    write_league(tmp_path)
    sheet = tmp_path / "Lag och GM.xlsx"
    assert main(["--source", str(tmp_path), "--out", str(sheet)]) == 0

    workbook = openpyxl.load_workbook(sheet)
    rows = {r[0]: r for r in workbook["Lag"].iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {
        "Alpha Wolves",
        "Beta Bears",
        "Epsilon Eagles",
        "Alpha Dogs",
        "Gamma Goats",
        "Delta Force",
    }
    assert rows["Alpha Dogs"][1] == "Anna Alfa"  # prefilled from `Namn och lag`
    assert "2022/23 Div 1 Gretzky" in rows["Alpha Dogs"][2]

    # The founder says Alpha Dogs was really Bertil's team.
    lag = workbook["Lag"]
    for row in lag.iter_rows(min_row=2):
        if row[0].value == "Alpha Dogs":
            row[1].value = "Bertil Beta"
    workbook.save(sheet)

    builder = LeagueBuilder(load(tmp_path, tmp_path / "none", tmp_path / "none"))
    league = builder.build()["league"]
    dogs = next(t for t in league["teams"] if t["name"] == "Alpha Dogs")
    assert dogs["managerId"] == "beta-bears"
