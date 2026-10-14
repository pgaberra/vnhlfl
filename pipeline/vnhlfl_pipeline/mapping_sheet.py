"""Writes `Lag och GM.xlsx`: every team name the league has used, for the founder to map to a GM.

    python -m vnhlfl_pipeline.mapping_sheet --source ../source-data \
        --out "../source-data/Lag och GM.xlsx"

Rows already mapped in `Namn och lag` come prefilled, so the founder only checks them and fills
the blanks. The filled-in file goes back in `source-data/` and takes precedence over
`Namn och lag` (see `sources.read_owners`). It contains real names: never commit it.
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .build import display_team, load
from .normalize import clean, key, season_label

SHEET = "Lag"
TEAM_COLUMN = "Lagnamn"
GM_COLUMN = "GM"

INSTRUCTIONS = [
    "Koppla varje lagnamn till rätt GM",
    "",
    "Fliken 'Lag' har ett lagnamn per rad: alla namn som förekommit i ligan sedan 2021/22.",
    "Fyll i kolumnen GM med deltagarens namn. Gula rader saknar GM.",
    "",
    "Så här:",
    "1. Samma person ska alltid skrivas exakt likadant. Har någon haft flera lagnamn genom åren",
    "   skriver du samma GM-namn på alla de raderna, så hänger karriären ihop på sidan.",
    "2. Rullistan i GM-kolumnen innehåller alla GM:s som redan finns. Välj därifrån när du kan.",
    "   Är det en ny deltagare skriver du namnet själv (Excel kan fråga om du vill behålla det).",
    "3. Rader som redan är ifyllda kommer från 'Namn och lag'. Rätta dem om något är fel.",
    "4. Kolumnen Förslag är bara en gissning utifrån liknande lagnamn. Kontrollera den.",
    "5. Byter ett lag GM mitt i säsongen syns det som två lagnamn. Skriv rätt GM på varje.",
    "",
    "Namnen publiceras inte: sidan visar bara lagnamn tills alla har sagt ja till annat.",
]


@dataclass
class _Seen:
    spellings: Counter[str] = field(default_factory=Counter)
    where: set[tuple[str, str]] = field(default_factory=set)


def team_rows(source: Path) -> list[dict[str, str]]:
    src = load(source, source / "no-curation", source / "no-transcribed")
    seen: dict[str, _Seen] = defaultdict(_Seen)
    appearances = [(r.season, r.division, r.team) for r in src.standings]
    appearances += [(p.season, p.division, p.team) for p in src.picks]
    for season, division, team in appearances:
        entry = seen[key(team)]
        entry.spellings[display_team(clean(team))] += 1
        entry.where.add((season, clean(division)))

    known = src.owners.by_team
    rows = []
    for team_key, entry in sorted(seen.items(), key=lambda kv: (min(kv[1].where), kv[0])):
        gm = known.get(team_key, "")
        rows.append(
            {
                "team": entry.spellings.most_common(1)[0][0],
                "gm": gm,
                "where": "; ".join(f"{season_label(s)} {d}" for s, d in sorted(entry.where)),
                "suggestion": "" if gm else suggest(team_key, known),
            }
        )
    return rows


def suggest(team_key: str, known: dict[str, str]) -> str:
    flat = team_key.replace(" ", "")
    best, score = "", 0.0
    for other in known:
        ratio = SequenceMatcher(None, flat, other.replace(" ", "")).ratio()
        if ratio > score:
            best, score = other, ratio
    return f"{known[best]}? (likt laget '{best}')" if score >= 0.85 else ""


def write(rows: list[dict[str, str]], gms: list[str], out: Path) -> None:
    workbook = openpyxl.Workbook()
    intro = workbook.active
    intro.title = "Instruktioner"
    for line in INSTRUCTIONS:
        intro.append([line])
    intro["A1"].font = Font(bold=True, size=14)
    intro.column_dimensions["A"].width = 100

    sheet = workbook.create_sheet(SHEET)
    headers = [TEAM_COLUMN, GM_COLUMN, "Förekommer (säsong, division)", "Förslag"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row["team"], row["gm"], row["where"], row["suggestion"]])
    last = sheet.max_row
    for column, width in zip("ABCD", (34, 30, 60, 46), strict=True):
        sheet.column_dimensions[column].width = width
    for cells in sheet.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{last}"
    yellow = PatternFill(start_color="FFF2B3", end_color="FFF2B3", fill_type="solid")
    sheet.conditional_formatting.add(
        f"A2:D{last}", FormulaRule(formula=["LEN(TRIM($B2))=0"], fill=yellow)
    )

    people = workbook.create_sheet("GM-lista")
    people.append(["GM:s som redan finns"])
    people["A1"].font = Font(bold=True)
    for name in gms:
        people.append([name])
    people.column_dimensions["A"].width = 34

    # A dropdown of known GMs that still accepts a new name (information style, not a stop).
    choice = DataValidation(
        type="list",
        formula1=f"='GM-lista'!$A$2:$A${len(gms) + 1}",
        allow_blank=True,
        errorStyle="information",
        errorTitle="Ny GM?",
        error="Namnet finns inte i listan. Klicka OK om det är en ny deltagare.",
    )
    sheet.add_data_validation(choice)
    choice.add(f"B2:B{last}")
    workbook.active = 1
    workbook.save(out)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="vnhlfl_pipeline.mapping_sheet", description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    rows = team_rows(args.source)
    gms = sorted({r["gm"] for r in rows if r["gm"]}, key=str.lower)
    write(rows, gms, args.out)
    missing = sum(1 for r in rows if not r["gm"])
    print(f"Wrote {args.out}: {len(rows)} team names, {missing} without a GM, {len(gms)} GMs")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
